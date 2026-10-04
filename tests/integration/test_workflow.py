"""End-to-end workflow through the service layer (mocked model, real document pipeline)."""

from __future__ import annotations

import asyncio
import hashlib
import io

import pytest
from docx import Document

from app.api.schemas import DecisionItem
from app.core.errors import DeliveryBlocked, InvalidState, LLMUnavailable, NotFound, StaleReview
from app.document.docx_parser import parse_docx
from app.resume.proposal_models import Decision
from tests.conftest import fixture_bytes
from tests.support.service_harness import JD, create_and_analyze, make_service, standard_batch

pytestmark = pytest.mark.renderer


def _decide(h, sid, token, revision, **decisions):
    items = [DecisionItem(edit_id=e, decision=Decision(d[0]), confirm_high_risk=d[1]) for e, d in decisions.items()]
    return h.service.decide(sid, token, revision, items)


@pytest.fixture
def data():
    return fixture_bytes("simple_one_page")


def test_full_workflow_upload_review_finalize_download(tmp_path, data):
    original_digest = hashlib.sha256(data).hexdigest()
    h = make_service(tmp_path, [standard_batch(data)])
    sid, token = create_and_analyze(h, data)

    view = h.service.get_session(sid, token)
    assert view.status == "awaiting_review"
    assert len(view.proposals) == 3 and view.counts["pending"] == 3
    assert view.analysis.alignment_estimate.coverage_percent is not None
    by_text = {p.original_text: p for p in view.proposals}
    migrate = by_text["Worked on the migration of"]
    lead = by_text["Responsible for code reviews and mentoring"]
    keyword = by_text["Built REST APIs in Python and FastAPI"]
    assert lead.requires_user_confirmation and not migrate.requires_user_confirmation
    assert any(seg.op == "insert" for seg in migrate.diff)
    assert migrate.paragraph_after.startswith("Migrated the billing system")

    # Accept one, explicitly confirm the high-risk one, reject the third.
    view = _decide(
        h,
        sid,
        token,
        view.review_revision,
        **{migrate.edit_id: ("accepted", False), lead.edit_id: ("accepted", True), keyword.edit_id: ("rejected", False)},
    )
    assert view.counts == {"total": 3, "pending": 0, "accepted": 2, "rejected": 1}

    result = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert result.decision == "PASS", result.report.reasons
    assert result.downloadable and result.is_current
    report = result.report
    assert {c.status for c in report.accepted_edit_verification} == {"applied_as_approved"}
    assert report.page_count.original == report.page_count.output == 1
    assert len(h.provider.requests) == 1  # finalization never calls the model

    link = h.service.issue_download_link(sid, token, result.finalization_id)
    payload = h.service.resolve_download(link.url.rsplit("/", 1)[1])
    assert payload.file_name == "Jordan Lee CV-edited.docx"
    texts = [p.text for p in Document(io.BytesIO(payload.data)).paragraphs]
    assert "Migrated the billing system to PostgreSQL." in texts
    assert "Led code reviews and mentored two junior engineers." in texts
    assert "Built REST APIs in Python and FastAPI serving 2M requests per day." in texts  # rejected edit absent
    # Derived from the uploaded document (same styles, same paragraph count), not a new template.
    before, after = parse_docx(data)[0], parse_docx(payload.data)[0]
    assert [b.style_id for b in before.blocks] == [b.style_id for b in after.blocks]
    assert "ATS-Optimized Resume" not in texts
    # Original immutable.
    assert hashlib.sha256(data).hexdigest() == original_digest
    stored = h.service.store.get(h.service.repo.find_artifact(sid, "original"))
    assert hashlib.sha256(stored).hexdigest() == original_digest

    events = [e["event"] for e in h.service.audit_trail(sid, token)]
    for expected in (
        "session_created",
        "proposal_created",
        "edit_accepted",
        "edit_rejected",
        "edit_applied",
        "finalization_completed",
        "downloaded",
    ):
        assert expected in events


def test_finalization_is_idempotent(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)])
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    migrate = next(p for p in view.proposals if p.original_text == "Worked on the migration of")
    view = _decide(h, sid, token, view.review_revision, **{migrate.edit_id: ("accepted", False)})
    first = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    second = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert first.finalization_id == second.finalization_id
    payload = h.service.resolve_download(h.service.issue_download_link(sid, token, first.finalization_id).url.rsplit("/", 1)[1])
    text = "\n".join(p.text for p in Document(io.BytesIO(payload.data)).paragraphs)
    assert text.count("Migrated") == 1


def test_high_risk_edit_requires_confirmation(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    lead = next(p for p in view.proposals if p.requires_user_confirmation)
    with pytest.raises(InvalidState) as exc:
        _decide(h, sid, token, view.review_revision, **{lead.edit_id: ("accepted", False)})
    assert exc.value.details["reason"] == "confirmation_required"


def test_stale_review_revision_rejected(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    eid = view.proposals[0].edit_id
    _decide(h, sid, token, view.review_revision, **{eid: ("accepted", False)})
    with pytest.raises(StaleReview):
        _decide(h, sid, token, view.review_revision, **{eid: ("rejected", False)})  # second tab with old revision
    with pytest.raises(StaleReview):
        asyncio.run(h.service.finalize(sid, token, view.review_revision))


def test_finalize_without_accepted_edits_rejected(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    with pytest.raises(InvalidState):
        asyncio.run(h.service.finalize(sid, token, view.review_revision))


def test_changed_decisions_invalidate_previous_download(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)])
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    by_text = {p.original_text: p for p in view.proposals}
    a, c = by_text["Worked on the migration of"], by_text["Built REST APIs in Python and FastAPI"]
    view = _decide(h, sid, token, view.review_revision, **{a.edit_id: ("accepted", False)})
    first = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    link = h.service.issue_download_link(sid, token, first.finalization_id)
    view = _decide(h, sid, token, view.review_revision, **{c.edit_id: ("accepted", False)})
    with pytest.raises(DeliveryBlocked):
        h.service.issue_download_link(sid, token, first.finalization_id)
    with pytest.raises(DeliveryBlocked):
        h.service.resolve_download(link.url.rsplit("/", 1)[1])  # already-issued link no longer valid either


def test_review_required_needs_acknowledgement(tmp_path):
    data = fixture_bytes("multipage")
    index, _ = parse_docx(data)
    from tests.support.llm_fakes import batch_json, proposal

    long_edit = proposal(
        index,
        "Delivered feature set 1 for the Northwind",
        "coordinating with design and QA",
        "coordinating closely with the design and QA teams each sprint",
        evidence=[
            (
                "Delivered feature set 1 for the Northwind",
                "coordinating with design and QA to ship on a two-week cadence and documenting decisions for future maintainers",
            )
        ],
    )
    h = make_service(tmp_path, [batch_json(long_edit)])
    sid, token = create_and_analyze(h, data, jd="Engineer with design and QA collaboration.")
    view = h.service.get_session(sid, token)
    assert len(view.proposals) == 1, view.analysis
    view = _decide(h, sid, token, view.review_revision, **{view.proposals[0].edit_id: ("accepted", False)})
    result = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert result.decision == "REVIEW_REQUIRED", result.report.reasons
    assert not result.downloadable and "acknowledge" in result.blocked_reason
    with pytest.raises(DeliveryBlocked):
        h.service.issue_download_link(sid, token, result.finalization_id)
    acknowledged = h.service.acknowledge(sid, token, result.finalization_id)
    assert acknowledged.downloadable
    h.service.issue_download_link(sid, token, result.finalization_id)


def test_failed_gate_cannot_be_bypassed(tmp_path, data, monkeypatch):
    """A corrupted output (simulated editor defect) is never downloadable, and retries return the same FAIL."""
    from app.workflow import nodes

    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    eid = next(p.edit_id for p in view.proposals if p.original_text == "Worked on the migration of")
    view = _decide(h, sid, token, view.review_revision, **{eid: ("accepted", False)})

    real_apply = nodes.apply_edits

    def corrupting_apply(original, index, ops):
        result = real_apply(original, index, ops)
        from app.evaluation.regressions import font_size_change

        result.output = font_size_change(result.output, index, {ops[0].location_id})
        return result

    monkeypatch.setattr(nodes, "apply_edits", corrupting_apply)
    result = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert result.decision == "FAIL" and not result.downloadable
    with pytest.raises(DeliveryBlocked):
        h.service.issue_download_link(sid, token, result.finalization_id)
    with pytest.raises(InvalidState):
        h.service.acknowledge(sid, token, result.finalization_id)  # FAIL cannot be acknowledged away
    monkeypatch.setattr(nodes, "apply_edits", real_apply)
    retry = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert retry.finalization_id == result.finalization_id and retry.decision == "FAIL"


def test_stale_source_edit_fails_safely(tmp_path, data):
    """If the stored proposal no longer matches the document, application fails and nothing is delivered."""
    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    eid = view.proposals[0].edit_id
    view = _decide(h, sid, token, view.review_revision, **{eid: ("accepted", False)})
    conn = h.service.repo._conn()
    payload = conn.execute("SELECT payload_json FROM proposals WHERE edit_id=?", (eid,)).fetchone()[0]
    conn.execute(
        "UPDATE proposals SET payload_json=? WHERE edit_id=?", (payload.replace('"original_text":"', '"original_text":"X', 1), eid)
    )
    result = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    assert result.decision == "FAIL"
    assert result.edit_failures and result.edit_failures[0]["code"] == "text_mismatch"


def test_analysis_failure_is_reported_and_retryable(tmp_path, data):
    h = make_service(tmp_path, [LLMUnavailable("The language model timed out."), standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    assert view.status == "analysis_failed"
    assert view.error.code == "llm_unavailable"
    assert view.analysis is None  # no fabricated score
    asyncio.run(h.service.run_analysis(sid))
    assert h.service.get_session(sid, token).status == "awaiting_review"


def test_analysis_runs_once_under_concurrency(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data)], renderer=False)
    created = h.service.create_session(
        data=data, file_name="cv.docx", content_type="", job_description=JD, company_details="", candidate_notes=""
    )

    async def both():
        await asyncio.gather(h.service.run_analysis(created.session_id), h.service.run_analysis(created.session_id))

    asyncio.run(both())
    assert len(h.provider.requests) == 1


def test_sessions_are_isolated(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data), standard_batch(data)], renderer=False)
    sid_a, token_a = create_and_analyze(h, data)
    sid_b, token_b = create_and_analyze(h, data)
    with pytest.raises(NotFound):
        h.service.get_session(sid_a, token_b)
    with pytest.raises(NotFound):
        h.service.get_session(sid_a, None)
    view_b = h.service.get_session(sid_b, token_b)
    with pytest.raises(NotFound):
        _decide(h, sid_a, token_b, 1, **{view_b.proposals[0].edit_id: ("accepted", False)})
    with pytest.raises(NotFound):
        h.service.get_session("../" + sid_a, token_a)


def test_download_token_tampering_and_expiry(tmp_path, data):
    from app.core.errors import Unauthorized
    from app.core.security import sign_download_token

    h = make_service(tmp_path, [standard_batch(data)])
    sid, token = create_and_analyze(h, data)
    view = h.service.get_session(sid, token)
    eid = next(p.edit_id for p in view.proposals if p.original_text == "Worked on the migration of")
    view = _decide(h, sid, token, view.review_revision, **{eid: ("accepted", False)})
    fin = asyncio.run(h.service.finalize(sid, token, view.review_revision))
    signed = h.service.issue_download_link(sid, token, fin.finalization_id).url.rsplit("/", 1)[1]
    with pytest.raises(Unauthorized):
        h.service.resolve_download(signed[:-2] + ("AA" if not signed.endswith("AA") else "BB"))
    expired, _ = sign_download_token(h.service.settings.signing_secret, sid, fin.finalization_id, 30, now=0)
    with pytest.raises(Unauthorized):
        h.service.resolve_download(expired)
    forged, _ = sign_download_token("another-secret-" + "y" * 32, sid, fin.finalization_id, 300)
    with pytest.raises(Unauthorized):
        h.service.resolve_download(forged)


def test_delete_and_retention(tmp_path, data):
    h = make_service(tmp_path, [standard_batch(data), standard_batch(data)], renderer=False)
    sid, token = create_and_analyze(h, data)
    session_dir = h.service.store.root / sid
    assert session_dir.exists()
    h.service.delete_session(sid, token)
    assert not session_dir.exists()
    with pytest.raises(NotFound):
        h.service.get_session(sid, token)

    sid2, _ = create_and_analyze(h, data)
    h.service.repo._conn().execute("UPDATE sessions SET expires_at='2000-01-01T00:00:00+00:00' WHERE id=?", (sid2,))
    assert h.service.purge_expired() == 1
    assert not (h.service.store.root / sid2).exists()
